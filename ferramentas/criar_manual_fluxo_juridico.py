from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "plugins" / "sa-fluxo-juridico" / "assets" / "papeis-timbrados" / "PAPEL-TIMBRADO-OFICIAL-SA-ADVOCACIA.docx"
OUTPUT = ROOT / "guias" / "MANUAL-DO-ADVOGADO-FLUXO-JURIDICO-SA-ADVOCACIA.docx"

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)
WHITE = RGBColor(255, 255, 255)
BEIGE = "D9C9BA"
LIGHT_BEIGE = "F7F3F0"
LIGHT_GRAY = "D9D9D9"


def set_run_font(run, name="Arial", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths, header=True):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if header and row_index == 0:
                set_cell_shading(cell, BEIGE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_BEIGE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=10.5, bold=(header and row_index == 0), color=BLACK)
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_paragraph(doc, text="", style=None, size=11, bold=False, color=BLACK, space_before=0, space_after=8, alignment=None, italic=False):
    p = doc.add_paragraph(style=style)
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True, color=BLACK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    marker = p.add_run("• ")
    set_run_font(marker, size=11, color=BLACK)
    run = p.add_run(text)
    set_run_font(run, size=11, color=BLACK)
    return p


def add_numbered(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(f"{number}. ")
    set_run_font(r1, size=11, bold=True)
    r2 = p.add_run(title)
    set_run_font(r2, size=11, bold=True)
    r3 = p.add_run(f": {text}")
    set_run_font(r3, size=11)
    return p


def add_prompt(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.right_indent = Cm(0.55)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, name="Courier New", size=10.5, color=BLACK)
    return p


def add_link(doc, url, label=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    relationship_id = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.text = label or url
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)
    return p


def preservar_timbrado_e_adicionar_descricao():
    """Restaura a arte do timbrado e adiciona texto alternativo sem mudança visual."""
    partes = ("word/header1.xml", "word/_rels/header1.xml.rels", "word/media/image1.png")
    with ZipFile(TEMPLATE, "r") as fonte, ZipFile(OUTPUT, "r") as destino:
        conteudos = {item.filename: destino.read(item.filename) for item in destino.infolist()}
        for parte in partes:
            conteudos[parte] = fonte.read(parte)

    cabecalho = conteudos["word/header1.xml"].decode("utf-8")
    cabecalho = cabecalho.replace(
        'name="Picture 1"',
        'name="Picture 1" descr="Marca Sousa Araújo Advocacia, Dra. Lidiane Sousa Araújo e contatos do escritório"',
        1,
    )
    conteudos["word/header1.xml"] = cabecalho.encode("utf-8")

    pacote = BytesIO()
    with ZipFile(OUTPUT, "r") as destino, ZipFile(pacote, "w", ZIP_DEFLATED) as novo:
        for item in destino.infolist():
            novo.writestr(item, conteudos[item.filename])
    OUTPUT.write_bytes(pacote.getvalue())


def main():
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"Modelo não encontrado: {TEMPLATE}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(TEMPLATE)
    clear_body(doc)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)

    title_style = doc.styles["Title"]
    title_style.font.color.rgb = BLACK
    title_properties = title_style.element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    add_paragraph(doc, "\u00a0", size=1, color=WHITE, space_after=46)
    add_paragraph(doc, "Manual de instalação do Fluxo Jurídico", style="Title", size=22, bold=True, space_after=4)
    add_paragraph(doc, "Como usar no ChatGPT Codex Claude e Claude Code", size=12, color=GRAY, space_after=16)
    add_paragraph(
        doc,
        "Este manual ensina como instalar e usar o pacote de trabalho jurídico da Sousa Araújo Advocacia. Você não precisa entender programação. Escolha abaixo a ferramenta que já utiliza e siga somente aquele caminho.",
        size=11,
        space_after=12,
    )

    add_heading(doc, "Link oficial para copiar")
    add_paragraph(doc, "Copie o endereço abaixo exatamente como está.", size=11, space_after=2)
    add_link(doc, "https://github.com/mix7producao/fluxo-juridico")

    add_heading(doc, "O que o pacote faz")
    for text in (
        "Organiza processos, documentos, conversas, áudios, vídeos e demais provas sem alterar os originais.",
        "Cria índice, linha do tempo, resumo, perguntas pendentes e identificação permanente das provas, como PROVA-0001.",
        "Produz e revisa peças jurídicas em Word, depois que o caso estiver organizado.",
        "Aplica o papel timbrado oficial, a formatação e as cores definidas pelo escritório.",
        "Ajuda a criar conteúdo jurídico institucional e educativo dentro das regras do escritório.",
    ):
        add_bullet(doc, text)

    add_heading(doc, "Qual opção escolher")
    add_bullet(doc, "ChatGPT: recomendado para a equipe. O administrador importa uma vez.")
    add_bullet(doc, "Codex: recomendado para organizar pastas, arquivos e peças em Word.")
    add_bullet(doc, "Claude: uso simples no navegador por meio de um Projeto conectado ao GitHub.")
    add_bullet(doc, "Claude Code: trabalho completo com pastas e arquivos, instalado com dois comandos.")

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_paragraph(doc, "\u00a0", size=1, color=WHITE, space_after=46)
    add_heading(doc, "ChatGPT para a equipe")
    add_paragraph(doc, "Este é o caminho recomendado para vários advogados. Somente o administrador faz a importação. Depois, cada advogado instala ou recebe o plug-in conforme a política escolhida pelo escritório.", size=11)
    add_numbered(doc, 1, "Abra a Administração", "entre em Plug-ins e escolha Adicionar e Importar marketplace.")
    add_numbered(doc, 2, "Cole o link oficial", "no campo Origem.")
    add_numbered(doc, 3, "Deixe Caminho vazio", "porque o arquivo do marketplace já está na raiz do repositório.")
    add_numbered(doc, 4, "Deixe Branch tag ou commit vazio", "para acompanhar a versão principal e receber as atualizações futuras.")
    add_numbered(doc, 5, "Clique em Importar marketplace", "e autorize o acesso ao GitHub quando aparecer.")
    add_numbered(doc, 6, "Abra o plug-in SA Fluxo Jurídico", "e escolha a política Instalado para os advogados que devem recebê-lo automaticamente.")
    add_numbered(doc, 7, "Abra um novo chat", "e peça a tarefa normalmente.")
    add_paragraph(doc, "Se você não encontrar o menu Administração, peça esta instalação ao administrador da conta do escritório.", size=10.5, bold=True, space_before=3)

    add_heading(doc, "Codex")
    add_paragraph(doc, "O Codex é a opção mais prática para organizar uma pasta inteira de caso, converter arquivos, transcrever materiais e produzir documentos. Se o plug-in já foi liberado no workspace do ChatGPT, abra Plug-ins, instale SA Fluxo Jurídico e inicie uma nova tarefa. No Codex CLI, digite o comando abaixo para abrir o catálogo de plug-ins.", size=11)
    add_prompt(doc, "/plugins")

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_paragraph(doc, "\u00a0", size=1, color=WHITE, space_after=46)
    add_heading(doc, "Claude no navegador")
    add_paragraph(doc, "No Claude comum, use um Projeto conectado ao GitHub. Ele recebe as instruções do escritório, mas não substitui o Codex ou o Claude Code para trabalhar diretamente com uma pasta inteira do computador.", size=11)
    add_numbered(doc, 1, "Crie um Projeto", "com o nome Fluxo Jurídico SA Advocacia.")
    add_numbered(doc, 2, "Na área de conhecimento do Projeto", "clique no sinal de mais e escolha GitHub.")
    add_numbered(doc, 3, "Cole o link oficial", "autorize o GitHub e selecione os arquivos do repositório.")
    add_numbered(doc, 4, "Quando o repositório for atualizado", "use Sincronizar antes de começar um trabalho novo.")

    add_heading(doc, "Claude Code")
    add_paragraph(doc, "Abra o Claude Code e envie os dois comandos abaixo, um de cada vez.", size=11)
    add_prompt(doc, "/plugin marketplace add mix7producao/fluxo-juridico")
    add_prompt(doc, "/plugin install sa-fluxo-juridico@sa-advocacia")
    add_paragraph(doc, "Para receber atualizações automáticas, abra /plugin, entre em Marketplaces, escolha sa-advocacia e ative Enable auto-update. Se aparecer a mensagem para recarregar, use /reload-plugins.", size=11)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_paragraph(doc, "\u00a0", size=1, color=WHITE, space_after=46)
    add_heading(doc, "Como começar depois da instalação")
    add_paragraph(doc, "Para organizar um caso, coloque todos os arquivos na pasta do caso e use a frase abaixo.", size=11, space_after=2)
    add_prompt(doc, "Quero abrir e organizar um novo caso jurídico.")
    add_paragraph(doc, "A ferramenta primeiro lembrará o que pode estar faltando. Quando você confirmar que terminou de enviar, ela preservará os originais e começará a organização.", size=11)
    add_paragraph(doc, "Depois que o caso estiver organizado, use a frase abaixo para produzir a peça.", size=11, space_after=2)
    add_prompt(doc, "Analise o caso organizado e produza a peça jurídica adequada em Word. Não invente informações e liste tudo que ainda depende de confirmação.")

    add_heading(doc, "Como a peça em Word será entregue")
    roles = doc.add_table(rows=1, cols=2)
    roles.rows[0].cells[0].text = "O pacote prepara"
    roles.rows[0].cells[1].text = "A advogada revisa"
    role_rows = (
        ("Arquivo editável em Word no papel timbrado oficial.", "Se o timbrado correto foi usado para aquele escritório."),
        ("Formatação, cores e estrutura definidas nas instruções do escritório.", "Estratégia, fatos, pedidos, valores, prazos e anexos."),
        ("Indicação clara do que está comprovado e do que precisa ser confirmado.", "Versão final antes de assinar ou protocolar."),
    )
    for left, right in role_rows:
        cells = roles.add_row().cells
        cells[0].text = left
        cells[1].text = right
    format_table(roles, (Cm(7.65), Cm(7.65)))

    add_heading(doc, "Regra de sigilo")
    add_paragraph(doc, "O GitHub guarda somente o manual, as instruções e os modelos vazios. Processos, conversas, provas, documentos pessoais e dados de clientes ficam exclusivamente na pasta de cada caso.", size=11, bold=True, space_after=0)

    doc.core_properties.title = "Manual de instalação do Fluxo Jurídico"
    doc.core_properties.subject = "Instalação e uso no ChatGPT, Codex, Claude e Claude Code"
    doc.core_properties.author = "Sousa Araújo Advocacia"
    doc.save(OUTPUT)
    preservar_timbrado_e_adicionar_descricao()
    print(OUTPUT)


if __name__ == "__main__":
    main()
